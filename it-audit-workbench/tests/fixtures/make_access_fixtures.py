"""Regenerate the access-test fixtures (needs openpyxl, used for fixtures only).

ad_users.csv     - Get-ADUser style export: US dates as PowerShell prints them.
hr_leavers.xlsx  - HR leavers report with real Excel date cells and a second sheet.

Planted cases (as at 2026-09-30, SLA 1 day, dormant after 90 days):
  bwong    leaver 2026-02-01, still enabled
  kpatel   leaver 2026-06-15, still enabled and logged on 2026-07-02
  tnguyen  leaver 2026-05-20, disabled 2026-06-05 (late)
  mrossi   leaver 2026-04-10, disabled same day (ok)
  lchen    leaver 2026-08-31, disabled 2026-09-01 (ok, within SLA)
  dkim     leaver 2026-09-29, enabled but within SLA on the as-at date (ok)
  fsilva   leaver 2026-10-15, after the as-at date (not yet due)
  ojones   leaver 2026-03-03, no account in the export (unmatched)
  rlee, pmartin dormant; anew never logged on
  svc-backup, admin.test, scanner01 generic accounts; training01 generic but disabled
"""
import csv
from datetime import datetime
from pathlib import Path

HERE = Path(__file__).resolve().parent
FIRST = ['james', 'maria', 'david', 'sarah', 'ahmed', 'li', 'peter', 'nina', 'omar', 'grace', 'ivan', 'yuki', 'carlos', 'emma',
         'ravi', 'zoe', 'hana', 'luke', 'sofia', 'noah', 'aisha', 'ben', 'chloe', 'dan', 'eva', 'finn', 'gina', 'hugo', 'iris',
         'jack', 'kira', 'leo', 'mia', 'nate', 'olga', 'paul', 'quinn', 'rosa', 'sam', 'tara']
LAST = ['smith', 'garcia', 'brown', 'khan', 'ali', 'wei', 'novak', 'berg', 'haddad', 'okafor', 'petrov', 'sato', 'diaz',
        'moore', 'iyer', 'clark', 'yamada', 'hall', 'costa', 'evans', 'malik', 'ford', 'wood', 'reid', 'lund', 'ward', 'ross',
        'shaw', 'kerr', 'lowe', 'cole', 'nash', 'page', 'rice', 'webb', 'york', 'zane', 'ames', 'bell', 'cruz']


def ps(d):
    """PowerShell en-US DateTime text, e.g. 9/28/2026 8:14:02 AM."""
    if not d:
        return ''
    dt = datetime.strptime(d, '%Y-%m-%d %H:%M') if ' ' in d else datetime.strptime(d, '%Y-%m-%d')
    return f"{dt.month}/{dt.day}/{dt.year} {dt.strftime('%I').lstrip('0')}:{dt.strftime('%M:%S %p')}"


def main():
    rows = []
    for i, (f, l) in enumerate(zip(FIRST, LAST)):
        sam = f[0] + l
        rows.append([sam, f'{f.title()} {l.title()}', 'True', ps(f'2026-09-{(i % 27) + 2:02d} 0{8 + i % 2}:{10 + i}'), ps('2026-01-15 10:00')])
    special = [
        ['bwong', 'Brian Wong', 'True', ps('2026-01-30 17:02'), ps('2025-11-02 09:00')],
        ['kpatel', 'Kiran Patel', 'True', ps('2026-07-02 22:41'), ps('2026-03-01 09:00')],
        ['tnguyen', 'Thao Nguyen', 'False', ps('2026-05-19 16:20'), ps('2026-06-05 11:12')],
        ['mrossi', 'Marco Rossi', 'False', ps('2026-04-09 15:00'), ps('2026-04-10 17:30')],
        ['lchen', 'Lin Chen', 'False', ps('2026-08-31 12:00'), ps('2026-09-01 08:05')],
        ['dkim', 'Daniel Kim', 'True', ps('2026-09-29 11:00'), ps('2026-02-01 09:00')],
        ['fsilva', 'Filipa Silva', 'True', ps('2026-09-29 16:00'), ps('2026-02-01 09:00')],
        ['rlee', 'Rachel Lee', 'True', ps('2026-05-01 09:30'), ps('2026-02-01 09:00')],
        ['pmartin', 'Paul Martin', 'True', ps('2026-06-25 14:10'), ps('2026-02-01 09:00')],
        ['anew', 'Alex New', 'True', '', ps('2026-02-01 09:00')],
        ['svc-backup', 'Backup Service', 'True', ps('2026-09-29 23:00'), ps('2025-06-01 09:00')],
        ['admin.test', 'Admin Test', 'True', ps('2026-09-01 10:00'), ps('2025-06-01 09:00')],
        ['scanner01', 'Floor 2 Scanner', 'True', ps('2026-09-28 07:00'), ps('2025-06-01 09:00')],
        ['training01', 'Training Account', 'False', ps('2025-12-01 10:00'), ps('2025-12-02 09:00')],
    ]
    rows[5:5] = special
    with open(HERE / 'ad_users.csv', 'w', newline='', encoding='utf-8') as fh:
        w = csv.writer(fh)
        w.writerow(['SamAccountName', 'DisplayName', 'Enabled', 'LastLogonDate', 'whenChanged'])
        w.writerows(rows)

    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'Leavers'
    ws.append(['Employee ID', 'Name', 'Network ID', 'Department', 'Termination Date'])
    leavers = [('E1001', 'Brian Wong', 'CORP\\bwong', 'Finance', '2026-02-01'),
               ('E1002', 'Kiran Patel', 'kpatel@acme.com', 'Sales', '2026-06-15'),
               ('E1003', 'Thao Nguyen', 'CORP\\TNguyen', 'IT', '2026-05-20'),
               ('E1004', 'Marco Rossi', 'CORP\\mrossi', 'HR', '2026-04-10'),
               ('E1005', 'Lin Chen', 'lchen@acme.com', 'Legal', '2026-08-31'),
               ('E1006', 'Daniel Kim', 'CORP\\dkim', 'Ops', '2026-09-29'),
               ('E1007', 'Filipa Silva', 'CORP\\fsilva', 'Ops', '2026-10-15'),
               ('E1008', 'Owen Jones', 'CORP\\ojones', 'Finance', '2026-03-03')]
    for e, n, nid, dep, d in leavers:
        ws.append([e, n, nid, dep, datetime.strptime(d, '%Y-%m-%d')])
        ws.cell(ws.max_row, 5).number_format = 'dd/mm/yyyy'
    notes = wb.create_sheet('Notes')
    notes.append(['Report run 30/09/2026 by HRIS'])
    wb.save(HERE / 'hr_leavers.xlsx')


if __name__ == '__main__':
    main()
